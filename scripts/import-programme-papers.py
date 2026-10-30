#!/usr/bin/env python3
"""Import public paper/session assignments and reconcile the conference programme.

Read the organizers' accepted-paper workbook, or a bounded connector snapshot with
--assignments-json. Only session labels and paper IDs are published; titles and
authors must agree with the existing accepted-paper publication before writing.
The input workbook and snapshot are never modified.
"""

import argparse
from collections import Counter
from copy import deepcopy
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
from zoneinfo import ZoneInfo

import openpyxl


ROOT = Path(__file__).resolve().parent.parent
SOURCE = "https://docs.google.com/spreadsheets/d/1eaf-WIvny0VXIeZrYAlUmVMcyHbpdbmIriB0rINRp08/edit"
ROOMS = ("Main Hall", "Room 1", "Room 2", "Room 3")


def clean(value):
    return " ".join(str(value).split()) if value is not None else ""


def number(value, label):
    if isinstance(value, bool):
        raise ValueError(f"Invalid {label}: {value!r}")
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    if re.fullmatch(r"[0-9]+", clean(value)):
        return int(clean(value))
    raise ValueError(f"Invalid {label}: {value!r}")


def session_number(value, kind):
    match = re.fullmatch(rf"{kind.title()} Session ([0-9]+)", clean(value))
    return number(match[1] if match else value, f"{kind} session number")


def read_workbook(path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    snapshot = {"oral": [], "poster": [], "oralSessions": [], "posterTitles": []}
    try:
        for kind, tab, header_row, headers in (
                ("oral", "oral presentation", 3,
                 ("paper id", "paper title", "authors", "oral session")),
                ("poster", "poster", 1,
                 ("paper id", "title", "authors", "poster session"))):
            sheet = workbook[tab]
            actual = tuple(clean(sheet.cell(header_row, column).value).casefold()
                           for column in range(2, 6))
            if actual != headers:
                raise ValueError(f"{tab}: unexpected B:E publication headers")
            for row_number, row in enumerate(sheet.iter_rows(
                    min_row=header_row + 1, min_col=2, max_col=5,
                    values_only=True), header_row + 1):
                if all(value in (None, "") for value in row):
                    continue
                paper_id, title, authors, session = row
                if paper_id in (None, ""):
                    raise ValueError(f"{tab}, row {row_number}: missing paper ID")
                snapshot[kind].append({
                    "id": number(paper_id, "paper ID"),
                    "title": clean(title),
                    "authors": clean(authors).replace(")*", ")"),
                    "session": session_number(session, kind),
                })
        sheet = workbook["Sessioni"]
        header = tuple(clean(sheet.cell(1, column).value).casefold()
                       for column in range(1, 4))
        if header != ("sessione", "titolo proposto", "paper"):
            raise ValueError("Sessioni: unexpected A:C headers")
        for row_number, row in enumerate(sheet.iter_rows(
                min_row=2, min_col=1, max_col=3, values_only=True), 2):
            session, title, papers = row
            if all(value in (None, "") for value in row):
                continue
            if clean(session).startswith("Poster Session "):
                if title not in (None, "") or papers not in (None, ""):
                    raise ValueError(f"Sessioni, row {row_number}: unexpected poster fields")
                snapshot["posterTitles"].append(clean(session))
                continue
            paper_ids = [number(value, "paper ID")
                         for value in clean(papers).split(",")]
            snapshot["oralSessions"].append({
                "number": number(session, "oral session number"),
                "title": clean(title), "paperIds": paper_ids,
            })
    finally:
        workbook.close()
    return snapshot


def build_assignments(snapshot, accepted, updated, source=SOURCE):
    date.fromisoformat(updated)
    groups = []
    oral_sessions = snapshot["oralSessions"]
    by_number = {number(item["number"], "oral session number"): item
                 for item in oral_sessions}
    if len(by_number) != len(oral_sessions) or set(by_number) != set(range(1, 17)):
        raise ValueError("Expected exactly the 16 numbered oral session themes")
    poster_titles = snapshot["posterTitles"]
    if len(poster_titles) != 3 or any(
            not clean(title).startswith(f"Poster Session {index} ")
            for index, title in enumerate(poster_titles, 1)):
        raise ValueError("Expected exactly the three numbered poster session themes")
    seen = set()
    for kind, expected_sessions in (("oral", range(1, 17)), ("poster", range(1, 4))):
        public = {paper["id"]: paper for group in accepted["groups"]
                  if group["id"] == kind for paper in group["papers"]}
        assigned = {session: [] for session in expected_sessions}
        rows = snapshot[kind]
        for row in rows:
            paper_id = number(row["id"], "paper ID")
            session = session_number(row["session"], kind)
            if paper_id in seen:
                raise ValueError(f"Duplicate paper ID: {paper_id}")
            if paper_id not in public:
                raise ValueError(f"Paper {paper_id} is not a published {kind} paper")
            if session not in assigned:
                raise ValueError(f"Paper {paper_id}: unknown {kind} session {session}")
            for field in ("title", "authors"):
                value = clean(row[field])
                if field == "authors":
                    value = value.replace(")*", ")")
                if value != clean(public[paper_id][field]):
                    raise ValueError(f"Paper {paper_id}: {field} differs from the publication")
            seen.add(paper_id)
            assigned[session].append(paper_id)
        if set(public) != {paper_id for ids in assigned.values() for paper_id in ids}:
            raise ValueError(f"The {kind} assignments must cover every published paper exactly once")
        for session, paper_ids in assigned.items():
            if kind == "oral":
                theme = by_number[session]
                if theme["paperIds"] != paper_ids:
                    raise ValueError(f"Oral session {session}: Sessioni paper order disagrees with assignments")
                title = f"Oral Session {session}: {clean(theme['title'])}"
                expected_count = 4 if session <= 4 else 3
            else:
                title = clean(poster_titles[session - 1])
                expected_count = 38
            if len(paper_ids) != expected_count or not title:
                raise ValueError(f"Unexpected paper count or title for {kind} session {session}")
            groups.append({"id": f"{kind}-{session}", "title": title, "paperIds": paper_ids})
    return {"updated": updated, "source": source, "groups": groups}


def oral_slots():
    """The source timetable's room order, with the existing identity of each slot."""
    slots = []
    for first, day, start, end, rooms, old_first in (
            (1, "2026-11-16", "11:40", "13:00", ROOMS, 1),
            (5, "2026-11-16", "15:00", "16:00", ROOMS, 5),
            (9, "2026-11-17", "10:00", "11:00", ROOMS, 5),
            (13, "2026-11-17", "14:30", "15:30",
             ("Main Hall", "Room 2", "Room 3", "Room 1"), 9)):
        for offset, room in enumerate(rooms):
            old_start = "11:00" if first == 1 else start
            identity = json.dumps([day, old_start, room, f"Oral Session {old_first + offset}"],
                                  ensure_ascii=False)
            identity_hash = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:8]
            slots.append({"number": first + offset, "date": day, "start": start,
                          "end": end, "room": room, "oldStart": old_start,
                          "id": f"session-{day}-{old_start.replace(':', '')}-{identity_hash}"})
    return slots


def apply_programme_papers(programme, assignments):
    """Apply the confirmed themes, counts and October 29 timetable correction."""
    result = deepcopy(programme)
    groups = {group["id"]: group for group in assignments["groups"]}
    expected = {f"oral-{index}" for index in range(1, 17)} | {
        f"poster-{index}" for index in range(1, 4)}
    if len(groups) != len(assignments["groups"]) or set(groups) != expected:
        raise ValueError("Programme paper assignments require 16 oral and three poster groups")
    for slot in oral_slots():
        candidates = [session for session in result["sessions"]
                      if session["type"] == "oral-session"
                      and session["date"] == slot["date"]
                      and session["room"] == slot["room"]
                      and session["start"] in (slot["start"], slot["oldStart"])]
        if len(candidates) != 1 or candidates[0]["end"] != slot["end"]:
            raise ValueError(f"Oral session {slot['number']}: timetable needs reconciliation")
        session = candidates[0]
        group = groups[f"oral-{slot['number']}"]
        session.update({"id": slot["id"], "start": slot["start"],
                        "title": group["title"], "paperIds": group["paperIds"],
                        "detail": f"{len(group['paperIds'])} presentations", "pending": False})
    if len([session for session in result["sessions"]
            if session["type"] == "oral-session"]) != 16:
        raise ValueError("Unexpected number of oral timetable sessions")
    for index in range(1, 4):
        candidates = [session for session in result["sessions"]
                      if re.match(rf"^Coffee Break \+ Poster Session {index}(?:$|\s|:)",
                                  session["title"])]
        if len(candidates) != 1:
            raise ValueError(f"Poster session {index}: timetable needs reconciliation")
        group = groups[f"poster-{index}"]
        session = candidates[0]
        old_title = f"Coffee Break + Poster Session {index}"
        identity = json.dumps([session["date"], session["start"], session["room"], old_title],
                              ensure_ascii=False)
        identity_hash = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:8]
        session.update({"id": f"session-{session['date']}-{session['start'].replace(':', '')}-{identity_hash}",
                              "title": f"Coffee Break + {group['title']}",
                              "paperIds": group["paperIds"],
                              "detail": f"{len(group['paperIds'])} posters"})
    barr = [session for session in result["sessions"]
            if "barr" in session["title"].casefold()]
    if len(barr) > 1:
        raise ValueError("Duplicate Michael Barr timetable session")
    if barr:
        session = barr[0]
        if tuple(session[key] for key in ("date", "start", "end", "room")) != (
                "2026-11-16", "11:00", "11:30", "Main Hall"):
            raise ValueError("Michael Barr timetable needs reconciliation")
    else:
        session = {"date": "2026-11-16", "start": "11:00", "end": "11:30",
                   "room": "Main Hall"}
        result["sessions"].append(session)
    session.update({"id": "keynote-michael-barr", "title": "Keynote: Michael Barr",
                    "type": "keynote", "detail": "Governor of the Federal Reserve.",
                    "pending": False, "href": ""})
    result["sessions"].sort(key=lambda item: (item["date"], item["start"]))
    if len({session["id"] for session in result["sessions"]}) != len(result["sessions"]):
        raise ValueError("Duplicate programme session ID after paper assignment")
    result["updated"] = max(result["updated"], assignments["updated"])
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", type=Path)
    parser.add_argument("--assignments-json", type=Path,
                        help="Read a bounded public-field connector snapshot instead of Excel")
    parser.add_argument("--updated", default=datetime.now(ZoneInfo("Europe/Rome")).date().isoformat())
    parser.add_argument("--source", default=SOURCE)
    parser.add_argument("--output", type=Path, default=ROOT / "data/programme-papers.json")
    parser.add_argument("--programme-output", type=Path, default=ROOT / "data/programme.json")
    args = parser.parse_args()
    if bool(args.workbook) == bool(args.assignments_json):
        parser.error("Supply one workbook or --assignments-json")
    snapshot = (json.loads(args.assignments_json.read_text(encoding="utf-8"))
                if args.assignments_json else read_workbook(args.workbook))
    accepted = json.loads((ROOT / "data/accepted-papers.json").read_text(encoding="utf-8"))
    assignments = build_assignments(snapshot, accepted, args.updated, args.source)
    programme = json.loads((ROOT / "data/programme.json").read_text(encoding="utf-8"))
    programme = apply_programme_papers(programme, assignments)
    for path, content in ((args.output, assignments), (args.programme_output, programme)):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(content, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.programme_output.resolve() == (ROOT / "data/programme.json").resolve():
        path = ROOT / "data/programme-announcements.json"
        announcements = json.loads(path.read_text(encoding="utf-8"))
        announcements = [item for item in announcements if item["id"] != "keynote-michael-barr"]
        path.write_text(json.dumps(announcements, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    counts = Counter(group["id"].split("-")[0] for group in assignments["groups"])
    print(f"Imported {counts['oral']} oral and {counts['poster']} poster groups; "
          f"{sum(len(group['paperIds']) for group in assignments['groups'])} assigned papers.")


if __name__ == "__main__":
    main()
