#!/usr/bin/env python3
"""Import the public ICAIF programme from the organizers' Excel workbook.

The workbook is read only. Notes and summary rows are never exported.
Requires openpyxl; see docs/programme-source.md for the refresh workflow.
"""

import argparse
from collections import Counter
from datetime import date, datetime, time
import hashlib
from html.parser import HTMLParser
import importlib.util
import json
from pathlib import Path
import re
from zoneinfo import ZoneInfo

import openpyxl


ROOT = Path(__file__).resolve().parent.parent
TIMEZONE = "Europe/Rome"
TYPES = {
    "registration", "workshop", "tutorial", "break", "industry",
    "competition", "opening", "keynote", "panel", "oral-session",
    "poster-session", "social", "banquet", "closing", "tbc",
}
DAY_SUBTITLES = {
    "2026-11-14": "Workshops & tutorials",
    "2026-11-15": "Workshops, tutorials & industry",
    "2026-11-16": "Main conference",
    "2026-11-17": "Main conference",
}

MORNING_PANEL_TITLE = (
    "From AI Scientific Discovery to Profitable Trading Strategies: "
    "What works, what transfers and what Falls Short?"
)
# Copy corrections keep the identity of the existing event and its saved links.
IDENTITY_TITLES = {
    "Industry day: AI Governance & DXC spotlight": "FINOS demo & panel",
    "Agostino Capponi Keynote & Perpetual Alpha Competition awards":
        "DXC session + Perpetual Alpha Competition awards",
    "QRT Startup Pitch Session": "QubeRT Startup Pitch Session",
    MORNING_PANEL_TITLE: "Panel discussion: Intesa, Domyn, Bloomberg & CFM",
    "Fireside chat": "Fireside chat: Raffaele Savi, Mauro Micillo & Stefano Pasquali",
}


def text(value):
    return " ".join(str(value).split()) if value is not None else ""


def parse_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for format_string in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text(value), format_string).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Invalid programme date: {value!r}")


def parse_time(value):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"Time is not expressed in whole minutes: {value!r}")
        return value.strftime("%H:%M")
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::00)?", text(value))
    if not match:
        raise ValueError(f"Invalid programme time: {value!r}")
    hour, minute = map(int, match.groups())
    try:
        return time(hour, minute).strftime("%H:%M")
    except ValueError as error:
        raise ValueError(f"Invalid programme time: {value!r}") from error


class WorkshopTitles(HTMLParser):
    """Read the site's existing workshop titles and anchors for checked links."""

    def __init__(self):
        super().__init__()
        self.anchor = ""
        self.reading = False
        self.fragments = []
        self.titles = {}

    def handle_starttag(self, tag, attributes):
        attributes = dict(attributes)
        if tag == "article":
            self.anchor = attributes.get("id", "")
        if tag == "h3" and attributes.get("data-index-label"):
            self.reading = True
            self.fragments = []

    def handle_data(self, value):
        if self.reading:
            self.fragments.append(value)

    def handle_endtag(self, tag):
        if tag == "h3" and self.reading:
            self.titles[self.anchor] = text("".join(self.fragments))
            self.reading = False


def title_key(value):
    # These punctuation differences do not change a workshop's identity.
    return text(value).replace("—", "-").replace("’", "'").rstrip(".").casefold()


def workshop_links():
    parser = WorkshopTitles()
    parser.feed((ROOT / "workshop.html").read_text(encoding="utf-8"))
    links = {title_key(title): f"/workshop/#{anchor}"
             for anchor, title in parser.titles.items() if anchor}
    security_title = parser.titles.get("financial-ai-security", "")
    short_security_title = "Financial AI Security, Privacy, and Safety"
    if security_title.startswith(short_security_title + ":"):
        links[title_key(short_security_title)] = "/workshop/#financial-ai-security"
    return links


def public_title(title, session_type):
    """Apply only the source-specific public copy decisions reviewed for this import."""
    detail = ""
    pending = False
    if not title:
        if session_type == "poster-session":
            return "Poster session", session_type, \
                "Title and presentation details to be confirmed.", True
        return "Session to be confirmed", "tbc", "", True
    if title.upper() == "TBD":
        return "Session to be confirmed", "tbc", "", True
    # The title explicitly identifies these formats even where the source Type
    # column is inconsistent. Keep the source's time, room, and title unchanged.
    if title == "Keynote Pasquali":
        session_type = "keynote"
    elif "poster session" in title.casefold() and session_type == "break":
        session_type = "poster-session"
    title = title.replace("The 3nd Workshop", "The 3rd Workshop")
    if title == "Financial AI Security, Privacy, and Safety:":
        title = title[:-1]
    elif title in ("Panel Discussion: Intesa, Domyn, Bloomberg, CFM, modera QRT",
                   "Panel discussion: Intesa, Domyn, Bloomberg & CFM", MORNING_PANEL_TITLE):
        title = MORNING_PANEL_TITLE
        detail = ("Participants: Andrea Prampolini (Intesa Sanpaolo), Stefano Pasquali (Domyn), "
                  "Anastasia Borovykh (CFM), Fabio Mercurio (Bloomberg).\n"
                  "Moderator: Georgios Papaioannou (QRT).")
    elif title in ("Panel Discussion: Banca D'Italia, Bundesbank, CEPR, ex FED tbd: ECB, BOE: modera Elena",
                   "AI, Systemic Risks and Central Banking"):
        title = "AI, Systemic Risks and Central Banking"
        session_type = "panel"
        detail = ("Participants: Kartik Anand (Bundesbank), Celso Brunetti "
                  "(Università di Roma Tor Vergata), Andrea Coletta (Bank of Italy), "
                  "and Jonas Fisher (Federal Reserve Bank of Chicago).\n"
                  "Moderator: Elena Carletti (Università Bocconi, Dean for Research).")
    elif title == "Banquet Dinner - Awards + speech Marc Mezard confermato":
        title = "Banquet dinner & awards"
        detail = "Speech by Marc Mézard."
    elif title == "ADIA Session":
        title = "ADIA Lab"
    elif title in ("Keynote: Magazzeni", "Keynote: Daniele Magazzeni"):
        title = "Keynote: Daniele Magazzeni"
        detail = "Group Chief AI Officer at UBS."
    elif title in ("Keynote Pasquali", "Keynote: Stefano Pasquali"):
        title = "Keynote: Stefano Pasquali"
        session_type = "keynote"
        detail = "Executive VP at Domyn."
    elif title in ("Fireside chat: Savi + Micillo + Pasquali",
                   "Fireside chat: Raffaele Savi, Mauro Micillo & Stefano Pasquali", "Fireside chat"):
        title = "Fireside chat"
        detail = ("Raffaele Savi (BlackRock Systematic), Mauro Micillo (Intesa Sanpaolo), "
                  "Stefano Pasquali (Domyn).")
    elif title in ("FINOS (demo+panel) 1h", "FINOS demo & panel"):
        title = "Industry day: AI Governance & DXC spotlight"
    elif title == "DXC session + Perpetual Alpha Competition awards":
        title = "Agostino Capponi Keynote & Perpetual Alpha Competition awards"
    elif title == "QubeRT Startup Pitch Session":
        title = "QRT Startup Pitch Session"
    return title, session_type, detail, pending


def import_workbook(input_path, updated):
    links = workshop_links()
    tutorial_data = json.loads((ROOT / "data/tutorials.json").read_text(encoding="utf-8"))
    tutorials = {tutorial["title"]: tutorial for tutorial in tutorial_data["tutorials"]}
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    sessions = []
    seen_days = set()
    for sheet in workbook:
        headers = {text(cell).casefold(): index
                   for index, cell in enumerate(next(sheet.iter_rows(values_only=True)))}
        required = {"date", "start", "end", "session title", "type", "room", "subsessions"}
        if not required <= headers.keys():
            raise ValueError(f"{sheet.title}: missing programme columns")
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            def value(column):
                index = headers[column]
                return row[index] if index < len(row) else None

            # Summary and internal-note rows have no date, start, or end.
            if all(value(column) in (None, "") for column in ("date", "start", "end")):
                continue
            try:
                session_date = parse_date(value("date"))
                start = parse_time(value("start"))
                end = parse_time(value("end"))
                if end <= start:
                    raise ValueError("Session must finish after it starts")
                if session_date not in DAY_SUBTITLES:
                    raise ValueError(f"Unexpected programme date: {session_date}")
                title = text(value("session title"))
                session_type = text(value("type")).casefold().replace(" ", "-")
                # One row has an explicit poster-session title and no type.
                if not session_type:
                    session_type = "poster-session" if "poster session" in title.casefold() else "tbc"
                if session_type not in TYPES:
                    raise ValueError(f"Unknown session type: {session_type}")
                title, session_type, detail, pending = public_title(title, session_type)
                # Claudio Tebaldi confirmed the 14:15 start on October 26.
                # Retain the workbook's 15:00 end; do not move the oral sessions.
                if title == "AI, Systemic Risks and Central Banking":
                    if (session_date, start, end, text(value("room"))) not in (
                            ("2026-11-16", "14:00", "15:00", "Main Hall"),
                            ("2026-11-16", "14:15", "15:00", "Main Hall")):
                        raise ValueError("Central banking panel schedule needs reconciliation "
                                         "with the October 26 organizer confirmation")
                    start = "14:15"
                room = text(value("room"))
                if not room:
                    room = "To be confirmed"
                    pending = True
                # Rachel Chung requested 16:30–18:00 on October 17 and confirmed
                # it on October 18. The earlier workbook block is cancelled.
                if (title == "Women in AI and Finance"
                        and (session_date, start, end, room)
                        == ("2026-11-14", "14:00", "16:00", "Room 2")):
                    continue
                # October 27: omit this placeholder until the organizers define it.
                if (title == "Session to be confirmed"
                        and (session_date, start, end, room)
                        == ("2026-11-15", "11:00", "13:00", "Room 3")):
                    continue
                count = value("subsessions")
                if isinstance(count, (int, float)) and count > 0:
                    if not float(count).is_integer():
                        raise ValueError(f"Non-integral presentation/poster count: {count}")
                    unit = "posters" if session_type == "poster-session" or "poster session" in title.casefold() else "presentations" if session_type == "oral-session" else ""
                    if unit:
                        detail = " ".join(part for part in (detail, f"{int(count)} {unit}") if part)
                identity = json.dumps([session_date, start, room, IDENTITY_TITLES.get(title, title)], ensure_ascii=False)
                identity_hash = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:8]
                href = links.get(title_key(title), "") if session_type == "workshop" else ""
                if session_type == "tutorial" and title in tutorials:
                    tutorial = tutorials[title]
                    actual = (session_date, start, end, room)
                    expected = tuple(tutorial[key] for key in ("date", "start", "end", "room"))
                    if actual != expected:
                        raise ValueError("Tutorial schedule differs from the supplied tutorial page data")
                    names = [presenter["name"] for presenter in tutorial["presenters"]]
                    joined = ", ".join(names[:-1]) + " and " + names[-1] if len(names) > 1 else names[0]
                    detail = " ".join(part for part in (detail, f"Presenters: {joined}.") if part)
                    href = f"/tutorials/#{tutorial['id']}"
                sessions.append({
                    "id": f"session-{session_date}-{start.replace(':', '')}-{identity_hash}",
                    "date": session_date,
                    "start": start,
                    "end": end,
                    "title": title,
                    "type": session_type,
                    "room": room,
                    "detail": detail,
                    "pending": pending,
                    "href": href,
                })
                seen_days.add(session_date)
            except ValueError as error:
                raise ValueError(f"{sheet.title}, row {row_number}: {error}") from error
    workbook.close()
    women_sessions = [session for session in sessions
                      if session["title"] == "Women in AI and Finance"]
    if (len(women_sessions) != 1
            or tuple(women_sessions[0][key] for key in ("date", "start", "end", "room"))
            != ("2026-11-14", "16:30", "18:00", "Room 2")):
        raise ValueError("Women in AI and Finance must retain only the confirmed "
                         "November 14, 16:30–18:00 Room 2 session")
    if seen_days != DAY_SUBTITLES.keys():
        raise ValueError("Workbook does not contain every conference day")
    if len({session["id"] for session in sessions}) != len(sessions):
        raise ValueError("Duplicate session IDs: check repeated date/start/room/title combinations")
    # Python's stable sort keeps source-row order within a simultaneous start.
    sessions.sort(key=lambda session: (session["date"], session["start"]))
    days = [{"date": day, "label": date.fromisoformat(day).strftime("%A"),
             "subtitle": DAY_SUBTITLES[day]} for day in sorted(seen_days)]
    programme = {"version": 1, "updated": updated, "timezone": TIMEZONE,
                 "days": days, "sessions": sessions}
    # The October 29 paper workbook supplies the session themes and assignments.
    # Its reviewed timetable correction also schedules Barr and moves the first
    # four oral sessions to 11:40, preserving their existing agenda identities.
    assignments_path = ROOT / "data/programme-papers.json"
    if assignments_path.exists():
        spec = importlib.util.spec_from_file_location(
            "programme_papers", ROOT / "scripts/import-programme-papers.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        assignments = json.loads(assignments_path.read_text(encoding="utf-8"))
        programme = module.apply_programme_papers(programme, assignments)
    return programme


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Path to an exported .xlsx programme workbook")
    parser.add_argument("--output", type=Path, default=ROOT / "data/programme.json")
    parser.add_argument("--updated", default=datetime.now(ZoneInfo(TIMEZONE)).date().isoformat(),
                        help="Public update date (YYYY-MM-DD; default: today in Rome)")
    arguments = parser.parse_args()
    date.fromisoformat(arguments.updated)
    data = import_workbook(arguments.input, arguments.updated)
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    counts = Counter(session["date"] for session in data["sessions"])
    print(f"Imported {len(data['sessions'])} sessions into {arguments.output}")
    for day in data["days"]:
        print(f"{day['date']}: {counts[day['date']]} sessions")


if __name__ == "__main__":
    main()
